import openpyxl
import sys
import os

# File paths
input_file = "docs/Rapport/Rapport Financier FATAPLUS - 2ém accompte.xlsx"
output_file = "docs/Rapport/Rapport Financier FATAPLUS - 2ém accompte_V2.xlsx"

if not os.path.exists(input_file):
    print(f"Error: Input file '{input_file}' not found.")
    sys.exit(1)

try:
    wb = openpyxl.load_workbook(input_file)
    print(f"Loaded workbook: {input_file}")
    
    # Try to find the relevant sheet
    sheet = wb.active
    print(f"Active sheet: {sheet.title}")
    
    # Check for "Frais bancaires"
    found = False
    last_row_index = 0
    
    # Iterate through rows to find where to insert or if it exists
    for row in sheet.iter_rows(min_row=1, max_col=5, values_only=False):
        cell_value = row[0].value
        if cell_value and isinstance(cell_value, str) and "frais bancaires" in cell_value.lower():
            found = True
            print("Found existing 'Frais bancaires' row.")
            # We could update the value here if we had specific numbers, 
            # but the task is to ensure it exists and includes all fees.
            # We'll append a note or ensure it's labeled correctly.
            row[0].value = "Frais bancaires (Tenue de compte, Virements, etc.)"
            break
        
        if cell_value:
             last_row_index = row[0].row

    if not found:
        print("Adding 'Frais bancaires' row...")
        # Simplistic insertion at the end or before totals if we could identify them.
        # For now, appending to the specific list of expenses often works best if no structure is strictly enforced,
        # but usually there is a table.
        # Let's try to find the "Total" row and insert before it.
        
        total_row = None
        for row in sheet.iter_rows(min_row=1, max_col=5, values_only=False):
            cell_value = row[0].value
            if cell_value and isinstance(cell_value, str) and "total" in cell_value.lower():
                total_row = row[0].row
                break
        
        if total_row:
            sheet.insert_rows(total_row)
            sheet.cell(row=total_row, column=1).value = "Frais bancaires (Tenue de compte, Virements, etc.)"
            sheet.cell(row=total_row, column=2).value = "100000" # Placeholder value based on previous context 100 000 Ar
            print(f"Inserted before Total row at {total_row}")
        else:
            # excessive fallback
            next_row = sheet.max_row + 1
            sheet.cell(row=next_row, column=1).value = "Frais bancaires (Tenue de compte, Virements, etc.)"
            sheet.cell(row=next_row, column=2).value = "100000"
            print("Appended at the end.")

    wb.save(output_file)
    print(f"Saved updated file to: {output_file}")

except Exception as e:
    print(f"An error occurred: {e}")
    sys.exit(1)
